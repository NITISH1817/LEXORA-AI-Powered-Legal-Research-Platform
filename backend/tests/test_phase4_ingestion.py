import pytest
import io
import os
import fitz
import docx
from fastapi.testclient import TestClient
from backend.app import app
from backend.database import SessionLocal
from backend.models.domain import Document, DocumentChunk, User
from backend.models.document_intel import DocumentPage, DocumentProcessingJob

client = TestClient(app)

@pytest.fixture(autouse=True)
def mock_embedding(monkeypatch):
    """Mock neural embedding generation for fast, deterministic ingestion tests."""
    monkeypatch.setattr("backend.services.vector_service.get_embedding", lambda text: [0.01] * 384)
    monkeypatch.setattr("backend.services.embedding_service.get_embedding", lambda text: [0.01] * 384)

# ---------------------------------------------------------------------------
# Synthetic Test Fixtures
# ---------------------------------------------------------------------------

def create_synthetic_pdf(pages_text: list) -> bytes:
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

def create_scanned_blank_pdf() -> bytes:
    doc = fitz.open()
    doc.new_page()  # 1 blank page with zero text
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

def create_synthetic_docx(paragraphs: list, table_data: list = None) -> bytes:
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = cell_value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.fixture(scope="module")
def user_tokens():
    """Create two researcher users and an admin user for cross-authorization tests."""
    tokens = {}
    users_data = [
        ("p4_researcher_a@lexora.ai", "pass123456", "Researcher A", "RESEARCHER"),
        ("p4_researcher_b@lexora.ai", "pass123456", "Researcher B", "RESEARCHER"),
        ("p4_admin@lexora.ai", "adminpass123", "Admin User", "ADMIN"),
    ]
    for email, password, full_name, role in users_data:
        client.post("/api/auth/register", json={
            "email": email,
            "password": password,
            "full_name": full_name,
            "invitation_code": "LEXORA-PILOT-2026-SUPERVISED"
        })
        login_resp = client.post("/api/auth/login", data={"username": email, "password": password})
        tokens[email] = login_resp.json()["access_token"]
        
        # Ensure role is set correctly in DB
        db = SessionLocal()
        u = db.query(User).filter(User.email == email).first()
        if u and u.role != role:
            u.role = role
            db.commit()
        db.close()
    return tokens

# ---------------------------------------------------------------------------
# 1. Successful PDF, DOCX, and TXT Ingestion & Extraction
# ---------------------------------------------------------------------------

def test_successful_pdf_ingestion(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = create_synthetic_pdf([
        "FACTS: The petitioner entered into an infrastructure concession agreement with the State.",
        "ISSUES: Whether termination of the concession agreement without prior notice violates natural justice."
    ])

    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("supreme_court_judgment.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    assert resp.status_code == 201
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    assert doc is not None
    assert doc.status == "READY"
    assert doc.page_count == 2
    assert len(doc.pages) == 2
    assert len(doc.chunks) >= 2
    
    # Verify page texts and sequence
    page1 = next(p for p in doc.pages if p.page_number == 1)
    assert "FACTS:" in page1.text
    page2 = next(p for p in doc.pages if p.page_number == 2)
    assert "ISSUES:" in page2.text
    db.close()

def test_successful_docx_ingestion(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = create_synthetic_docx(
        paragraphs=["ORDER: The interlocutory application for stay is hereby dismissed."],
        table_data=[["Item", "Claim Amount"], ["Liquidated Damages", "5000000"]]
    )

    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("court_order.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers
    )
    assert resp.status_code == 201
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    assert doc.status == "READY"
    assert doc.page_count >= 1
    # Verify table data was extracted alongside paragraphs
    full_text = "\n".join(p.text for p in doc.pages)
    assert "Liquidated Damages" in full_text
    assert "ORDER:" in full_text
    db.close()

def test_successful_txt_ingestion(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    txt_content = ("JUDGMENT: In the High Court of Delhi.\n" * 15).encode("utf-8")
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("brief_notes.txt", txt_content, "text/plain")},
        headers=headers
    )
    assert resp.status_code == 201
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    assert doc.status == "READY"
    assert len(doc.pages) >= 1
    assert "In the High Court of Delhi" in doc.pages[0].text
    db.close()

# ---------------------------------------------------------------------------
# 2. Content-Hash Duplicate Detection
# ---------------------------------------------------------------------------

def test_content_hash_duplicate_detection(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = create_synthetic_pdf(["Unique test authority content for deduplication test."])

    # First upload
    resp1 = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("original_case.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    assert resp1.status_code == 201
    doc1_id = resp1.json()["document_id"]

    # Second upload with identical bytes
    resp2 = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("duplicate_case.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    assert resp2.status_code == 201
    data2 = resp2.json()
    assert data2["document_id"] == doc1_id
    assert data2["duplicate_of"] == doc1_id
    assert "already exists" in data2["message"]

# ---------------------------------------------------------------------------
# 3. Invalid, Empty, Corrupt, and Unsupported Files
# ---------------------------------------------------------------------------

def test_empty_file_rejected(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.post(
        "/api/documents/upload",
        files={"file": ("empty.pdf", b"", "application/pdf")},
        headers=headers
    )
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()

def test_corrupted_pdf_rejected(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    corrupt_bytes = b"%PDF-invalid-corrupt-data-not-a-valid-structure"
    resp = client.post(
        "/api/documents/upload",
        files={"file": ("corrupt.pdf", corrupt_bytes, "application/pdf")},
        headers=headers
    )
    assert resp.status_code == 422
    assert "corrupted" in resp.json()["detail"].lower()

def test_corrupted_docx_rejected(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    corrupt_docx = b"PK\x03\x04corrupted_docx_without_word_document_xml"
    resp = client.post(
        "/api/documents/upload",
        files={"file": ("corrupt.docx", corrupt_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers
    )
    assert resp.status_code in (415, 422)

def test_unsupported_file_extension_rejected(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.post(
        "/api/documents/upload",
        files={"file": ("script.py", b"print('hello')", "text/x-python")},
        headers=headers
    )
    assert resp.status_code == 415
    assert "Unsupported file type" in resp.json()["detail"]

# ---------------------------------------------------------------------------
# 4. Page and Chunk Metadata Preservation
# ---------------------------------------------------------------------------

def test_chunk_page_metadata_preservation(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = create_synthetic_pdf([
        "FACTS: Paragraph on page one regarding the contract formation and breach.",
        "REASONING: Paragraph on page two analyzing the precedents on liquidated damages."
    ])
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("multipage_metadata_doc.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
    assert len(chunks) >= 2
    
    # Check that each chunk is associated with the correct document ID and page number
    for c in chunks:
        assert c.document_id == doc_id
        assert c.page_number in (1, 2)
        assert c.chunk_sequence is not None
        assert c.chunk_id_str is not None
        assert len(c.text) > 0
    db.close()

# ---------------------------------------------------------------------------
# 5. Scanned PDF & OCR-Required Behavior
# ---------------------------------------------------------------------------

def test_scanned_pdf_reports_ocr_required(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    blank_pdf = create_scanned_blank_pdf()
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("scanned_blank_case.pdf", blank_pdf, "application/pdf")},
        headers=headers
    )
    assert resp.status_code == 201
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    # Must report OCR_REQUIRED honestly when OCR is not available/configured
    assert doc.status == "OCR_REQUIRED"
    assert doc.ocr_required is True
    # Must NOT be marked as READY
    assert doc.status != "READY"
    db.close()

# ---------------------------------------------------------------------------
# 6. Processing State, Failure, and Retry Behavior
# ---------------------------------------------------------------------------

def test_document_retry_behavior(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    # Upload doc
    txt_content = b"Content for retry test."
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("retry_doc.txt", txt_content, "text/plain")},
        headers=headers
    )
    doc_id = resp.json()["document_id"]

    # Document is now READY; retrying a READY document reports already processed
    retry_resp = client.post(f"/api/documents/{doc_id}/retry", headers=headers)
    assert retry_resp.status_code == 200
    assert "already ready" in retry_resp.json()["message"]

    # Simulate FAILED status
    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    doc.status = "FAILED"
    db.commit()
    db.close()

    # Retrying a FAILED document succeeds
    retry_resp2 = client.post(f"/api/documents/{doc_id}/retry", headers=headers)
    assert retry_resp2.status_code == 200
    assert "Retry started" in retry_resp2.json()["message"]

# ---------------------------------------------------------------------------
# 7. Document Authorization & Cross-User Security
# ---------------------------------------------------------------------------

def test_cross_user_document_isolation(user_tokens):
    token_a = user_tokens["p4_researcher_a@lexora.ai"]
    token_b = user_tokens["p4_researcher_b@lexora.ai"]
    token_admin = user_tokens["p4_admin@lexora.ai"]

    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}
    headers_admin = {"Authorization": f"Bearer {token_admin}"}

    # User A uploads a private document
    txt_content = b"Confidential client privilege text for User A."
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("user_a_private.txt", txt_content, "text/plain")},
        headers=headers_a
    )
    doc_a_id = resp.json()["document_id"]

    # User B attempts to access User A's document -> 404 (does not disclose existence)
    resp_b = client.get(f"/api/documents/{doc_a_id}", headers=headers_b)
    assert resp_b.status_code == 404
    assert resp_b.json()["detail"] == "Document not found."

    # User B attempts to delete User A's document -> 404
    resp_b_del = client.delete(f"/api/documents/{doc_a_id}", headers=headers_b)
    assert resp_b_del.status_code == 404

    # Admin CAN access User A's document
    resp_admin = client.get(f"/api/documents/{doc_a_id}", headers=headers_admin)
    assert resp_admin.status_code == 200
    assert resp_admin.json()["id"] == doc_a_id

# ---------------------------------------------------------------------------
# 8. Document Reader and Source Retrieval API Contracts
# ---------------------------------------------------------------------------

def test_document_reader_endpoints(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = create_synthetic_pdf([
        "JUDGMENT: Page 1 content regarding statutory interpretation.",
        "ORDER: Page 2 content directing relief to the petitioner."
    ])
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("reader_contract_test.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    doc_id = resp.json()["document_id"]

    # 1. Document Detail: GET /api/documents/{id}
    detail_resp = client.get(f"/api/documents/{doc_id}", headers=headers)
    assert detail_resp.status_code == 200
    detail = detail_resp.json()
    assert detail["id"] == doc_id
    assert detail["page_count"] == 2
    assert "sections" in detail
    assert "metadata_fields" in detail
    assert "processing_logs" in detail

    # 2. Document Status: GET /api/documents/{id}/status
    status_resp = client.get(f"/api/documents/{doc_id}/status", headers=headers)
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] == "READY"
    assert status_resp.json()["progress"] == 100

    # 3. Source Page Retrieval: GET /api/documents/{id}/pages/{page}
    page_resp = client.get(f"/api/documents/{doc_id}/pages/1", headers=headers)
    assert page_resp.status_code == 200
    pdata = page_resp.json()
    assert pdata["page_number"] == 1
    assert "statutory interpretation" in pdata["text"]

    # 4. Invalid page number (e.g. 0) -> 400
    bad_page_resp = client.get(f"/api/documents/{doc_id}/pages/0", headers=headers)
    assert bad_page_resp.status_code == 400

    # 5. In-Document Search: GET /api/documents/{id}/search?q=statutory
    search_resp = client.get(f"/api/documents/{doc_id}/search?q=statutory", headers=headers)
    assert search_resp.status_code == 200
    hits = search_resp.json()
    assert len(hits) >= 1
    assert any("statutory" in h["text_snippet"].lower() for h in hits)

    # 6. Processing Logs: GET /api/documents/{id}/logs
    logs_resp = client.get(f"/api/documents/{doc_id}/logs", headers=headers)
    assert logs_resp.status_code == 200
    logs = logs_resp.json()
    assert isinstance(logs, list)
    assert len(logs) > 0
    assert any(l["stage"] == "EXTRACTION" for l in logs)

# ---------------------------------------------------------------------------
# 9. Document Deletion & Associated Data Cleanup
# ---------------------------------------------------------------------------

def test_document_deletion_cleanup(user_tokens):
    token = user_tokens["p4_researcher_a@lexora.ai"]
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = create_synthetic_pdf(["Content for deletion verification."])
    resp = client.post(
        "/api/documents/upload?sync=true",
        files={"file": ("to_be_deleted.pdf", pdf_bytes, "application/pdf")},
        headers=headers
    )
    doc_id = resp.json()["document_id"]

    db = SessionLocal()
    doc = db.query(Document).filter(Document.id == doc_id).first()
    storage_path = doc.storage_path
    assert os.path.exists(storage_path)
    db.close()

    # Delete document
    del_resp = client.delete(f"/api/documents/{doc_id}", headers=headers)
    assert del_resp.status_code == 200
    assert del_resp.json()["message"] == "Document deleted."

    # Verify physical file removed
    assert not os.path.exists(storage_path)

    # Verify DB records deleted
    db = SessionLocal()
    assert db.query(Document).filter(Document.id == doc_id).first() is None
    assert db.query(DocumentPage).filter(DocumentPage.document_id == doc_id).count() == 0
    assert db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).count() == 0
    assert db.query(DocumentProcessingJob).filter(DocumentProcessingJob.document_id == doc_id).first() is None
    db.close()
